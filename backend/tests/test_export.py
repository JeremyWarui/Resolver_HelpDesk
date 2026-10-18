"""The Excel and PDF reports: who may download them, and that they say what
`/analytics/` says for the same role and period."""

import io
from datetime import timedelta

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.tickets.lifecycle import assign_ticket, rate_ticket, transition_status
from tests import factories

pytestmark = pytest.mark.django_db

XLSX = "analytics:export-xlsx"
PDF = "analytics:export-pdf"


@pytest.fixture
def worked(
    nrb_electrical_ticket, msa_electrical_ticket, nrb_plumbing_ticket,
    nrb_electrician, msa_electrician, nrb_hos, msa_hos, requester,
):
    """One closed and rated NRB job, one MSA job on hold, one open NRB job."""
    assign_ticket(nrb_electrical_ticket, nrb_electrician, nrb_hos)
    for status in ("in_progress", "resolved"):
        transition_status(nrb_electrical_ticket, status, nrb_electrician)
    rate_ticket(nrb_electrical_ticket, requester, 4, "Good")
    assign_ticket(msa_electrical_ticket, msa_electrician, msa_hos)
    transition_status(msa_electrical_ticket, "in_progress", msa_electrician)
    transition_status(msa_electrical_ticket, "pending", msa_electrician, pending_reason="material_not_in_store")
    return nrb_electrical_ticket, msa_electrical_ticket, nrb_plumbing_ticket


def _workbook(api, user, **params):
    api.force_authenticate(user)
    response = api.get(reverse(XLSX), params)
    assert response.status_code == 200, response.content[:200]
    assert response["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml"
    )
    assert response["Content-Disposition"].startswith('attachment; filename="helpdesk-report-')
    return load_workbook(io.BytesIO(response.content))


def _summary(wb):
    return {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]}


def _sheet(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


# ─── Who may download what ──────────────────────────────────────────────────


@pytest.mark.parametrize(
    "role", ["admin_user", "manager", "nrb_hod", "nrb_hos", "nrb_electrician"]
)
def test_every_role_with_reports_downloads_the_workbook(api, request, worked, role):
    _workbook(api, request.getfixturevalue(role))


@pytest.mark.parametrize("role", ["manager", "nrb_hod", "nrb_hos", "nrb_electrician"])
def test_the_manager_hod_hos_and_technician_download_the_pdf(api, request, worked, role):
    api.force_authenticate(request.getfixturevalue(role))
    response = api.get(reverse(PDF))
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert response["Content-Disposition"].endswith('.pdf"')


def test_a_requester_downloads_neither_file(api, requester, worked):
    api.force_authenticate(requester)
    assert api.get(reverse(XLSX)).status_code == 403
    assert api.get(reverse(PDF)).status_code == 403


def test_the_downloads_need_a_login(api):
    assert api.get(reverse(XLSX)).status_code == 401
    assert api.get(reverse(PDF)).status_code == 401


# ─── What each workbook holds ───────────────────────────────────────────────


EVERY_CAMPUS_SHEETS = [
    "Summary", "Tickets raised", "Facility Health", "Campuses", "Trades", "Trades by campus",
    "Holds", "Tickets",
]
ONE_CAMPUS_SHEETS = ["Summary", "Tickets raised", "Facility Health", "Technicians", "Trades", "Holds", "Tickets"]


@pytest.mark.parametrize(
    "role, sheets, campuses, covers",
    [
        ("manager", EVERY_CAMPUS_SHEETS, {"Nairobi", "Mombasa"}, "Every campus"),
        ("admin_user", EVERY_CAMPUS_SHEETS, {"Nairobi", "Mombasa"}, "Every campus"),
        ("nrb_hos", ONE_CAMPUS_SHEETS, {"Nairobi"}, "Nairobi campus"),
    ],
)
def test_a_workbook_is_its_reports_page_over_the_roles_campuses(
    api, request, worked, role, sheets, campuses, covers
):
    wb = _workbook(api, request.getfixturevalue(role))
    assert wb.sheetnames == sheets
    assert {t["Campus"] for t in _sheet(wb, "Tickets")} == campuses
    assert _summary(wb)["Covers"] == covers


def test_a_technician_workbook_is_their_own_jobs_without_colleagues(
    api, nrb, electrical, nrb_electrician, requester, worked
):
    colleague = factories.make_technician("elec_two", nrb, [electrical])
    factories.make_ticket(requester, nrb, electrical, assigned_to=colleague)

    wb = _workbook(api, nrb_electrician)
    assert "By technician" not in wb.sheetnames and "By campus" not in wb.sheetnames
    assert {t["Assigned to"] for t in _sheet(wb, "Tickets")} == {
        f"{nrb_electrician.first_name} {nrb_electrician.last_name}".strip()
        or nrb_electrician.username
    }


def _pct(value):
    """A percentage as the workbook stores it: a fraction."""
    return None if value is None else pytest.approx(value / 100)


def test_the_workbook_says_what_the_reports_page_says(api, nrb_electrician, worked):
    """A technician's workbook and their reports page: same period, same numbers."""
    api.force_authenticate(nrb_electrician)
    page = api.get(reverse("analytics:analytics"), {"days": 30, "group_by": "trade"}).json()
    totals = page["totals"]
    wb = _workbook(api, nrb_electrician, days=30)
    summary = _summary(wb)

    assert summary["Jobs assigned to me"] == totals["total"]
    assert summary["Active now"] == totals["open_backlog"]
    assert summary["Overdue now"] == totals["overdue"]
    assert summary["Resolved on time"] == _pct(totals["resolution_sla_pct"])
    assert summary["Average rating (of 5)"] == totals["rating_avg"]
    assert summary["Jobs finished"] == totals["settled"]
    if totals["resolution_seconds_avg"] is not None:
        assert summary["Time to resolve"] == timedelta(
            seconds=totals["resolution_seconds_avg"]
        )

    rows = list(wb["By trade"].iter_rows(values_only=True))
    sheet = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
    assert set(sheet) == {row["label"] for row in page["breakdown"]}
    for row in page["breakdown"]:
        assert sheet[row["label"]]["Raised"] == row["total"]
        assert sheet[row["label"]]["Open"] == row["open"]
        assert sheet[row["label"]]["Completion %"] == _pct(row["completion_pct"])


def test_the_tickets_sheet_carries_the_rating_and_sla_outcome(api, manager, worked):
    nrb_job, _, _ = worked
    rows = {t["Ticket"]: t for t in _sheet(_workbook(api, manager), "Tickets")}

    row = rows[nrb_job.ticket_no]
    assert (row["Status"], row["Rating"], row["Rating comment"]) == ("Closed", 4, "Good")
    assert row["SLA"] == "Met"


def test_the_period_narrows_the_workbook(api, manager, worked):
    old = factories.age_ticket(worked[2], days=60)
    numbers = {t["Ticket"] for t in _sheet(_workbook(api, manager, days=30), "Tickets")}
    assert old.ticket_no not in numbers
    numbers = {t["Ticket"] for t in _sheet(_workbook(api, manager, days=90), "Tickets")}
    assert old.ticket_no in numbers


def test_an_empty_period_still_builds_both_formats(api, manager, nrb_hos):
    wb = _workbook(api, manager)
    assert _sheet(wb, "Tickets") == []
    api.force_authenticate(nrb_hos)
    assert api.get(reverse(PDF)).content.startswith(b"%PDF")


def test_a_bad_date_falls_back_to_the_default_period(api, manager, worked):
    _workbook(api, manager, date_from="not-a-date", date_to="2026-13-45")


# ─── Text people typed ──────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "typed", ['=HYPERLINK("http://evil","Click")', "+1+1", "-2", "@SUM(A1)"]
)
def test_a_formula_typed_by_a_requester_stays_text(api, manager, worked, typed):
    nrb_job, _, _ = worked
    type(nrb_job).objects.filter(pk=nrb_job.pk).update(
        location_detail=typed, rating_comment=typed
    )
    api.force_authenticate(manager)
    response = api.get(reverse(XLSX))
    ws = load_workbook(io.BytesIO(response.content))["Tickets"]

    headings = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == nrb_job.ticket_no)
    for column in ("Where", "Rating comment"):
        cell = row[headings.index(column)]
        assert cell.data_type == "s", column
        assert cell.value == "'" + typed


def test_control_characters_do_not_break_the_download(api, manager, worked):
    nrb_job, _, _ = worked
    type(nrb_job).objects.filter(pk=nrb_job.pk).update(
        location_detail="Room\x0b 4\x1b", rating_comment="ok\x07 thanks"
    )
    rows = {t["Ticket"]: t for t in _sheet(_workbook(api, manager), "Tickets")}
    assert rows[nrb_job.ticket_no]["Where"] == "Room 4"
    assert rows[nrb_job.ticket_no]["Rating comment"] == "ok thanks"


def test_a_campus_name_with_markup_or_accents_still_builds_the_pdf(api, nrb, nrb_hos):
    nrb.name = "Ńairobi <Main> & Annex"
    nrb.save(update_fields=["name"])
    api.force_authenticate(nrb_hos)
    response = api.get(reverse(PDF))

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    disposition = response["Content-Disposition"]
    assert disposition.startswith('attachment; filename="helpdesk-report-')
    assert disposition.isascii()


def test_an_hos_without_a_campus_is_labelled_plainly(nrb_hos):
    from apps.accounts.models import RoleAssignment
    from apps.analytics.export import scope_label

    RoleAssignment.objects.filter(user=nrb_hos).update(campus=None)
    assert scope_label(nrb_hos, "hos") == "Your campus"


def test_a_huge_period_is_clamped_not_a_server_error(api, manager, worked):
    _workbook(api, manager, days=999999)
    api.force_authenticate(manager)
    assert api.get(reverse("analytics:analytics"), {"days": 999999}).status_code == 200


# ─── The SLA outcome on the Tickets sheet ───────────────────────────────────


def test_the_sla_column_uses_the_one_definition_of_late(api, manager, worked):
    """A settled ticket with no target is unmeasured, as in every percentage."""
    nrb_job, _, _ = worked
    type(nrb_job).objects.filter(pk=nrb_job.pk).update(resolution_due_at=None)
    rows = {t["Ticket"]: t for t in _sheet(_workbook(api, manager), "Tickets")}
    assert rows[nrb_job.ticket_no]["SLA"] is None


def test_sla_outcome_matches_the_query_definitions(nrb_electrical_ticket):
    from apps.tickets.sla import sla_outcome

    t = nrb_electrical_ticket
    now = timezone.now()
    t.status, t.resolution_due_at, t.resolved_at = "closed", now, now - timedelta(hours=1)
    assert sla_outcome(t) == "met"
    t.resolved_at = now + timedelta(hours=1)
    assert sla_outcome(t) == "missed"
    t.resolved_at = None
    assert sla_outcome(t) == "missed"
    t.resolution_due_at = None
    assert sla_outcome(t) is None
    t.status, t.resolution_due_at = "in_progress", now - timedelta(hours=1)
    assert sla_outcome(t, now) == "overdue"
    t.status = "pending"
    assert sla_outcome(t, now) is None


def test_trend_days_are_kenyan_days(api, manager, nrb, electrical, requester, priorities):
    """01:00 in Nairobi is still the previous day in UTC; it counts as today."""
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from apps.tickets.models import Ticket

    ticket = factories.make_ticket(requester, nrb, electrical)
    nairobi_1am = datetime(2026, 10, 10, 1, 0, tzinfo=ZoneInfo("Africa/Nairobi"))
    Ticket.objects.filter(pk=ticket.pk).update(created_at=nairobi_1am)

    api.force_authenticate(manager)
    trend = api.get(
        reverse("analytics:analytics"),
        {"date_from": "2026-10-08T00:00:00+03:00", "date_to": "2026-10-12T00:00:00+03:00"},
    ).json()["trend"]
    assert [p["bucket"][:10] for p in trend] == ["2026-10-10"]


# ─── Other figures on /analytics/ ───────────────────────────────────────────


def test_analytics_reports_completion_escalation_split_and_rated_share(
    api, manager, worked
):
    nrb_job, msa_job, open_job = worked
    msa_job.current_level = "hod"
    msa_job.save(update_fields=["current_level"])
    open_job.current_level = "hos"
    open_job.save(update_fields=["current_level"])

    api.force_authenticate(manager)
    totals = api.get(reverse("analytics:analytics")).json()["totals"]

    assert totals["completion_pct"] == pytest.approx(100 / 3, abs=0.1)
    assert (totals["escalated_hos"], totals["escalated_hod"]) == (1, 1)
    assert (totals["closed"], totals["rated_pct"]) == (1, 100.0)


def test_a_technicians_headline_is_about_their_own_work(api, nrb_electrician, worked):
    """No figure about assigning work, which is the HOS's and never theirs."""
    summary = _summary(_workbook(api, nrb_electrician, days=30))

    assert "Jobs assigned to me" in summary and "Time to resolve" in summary
    for supervisors_only in ("Open (not yet assigned)", "Time to respond", "Responded on time"):
        assert supervisors_only not in summary


@pytest.mark.parametrize("narrowed", [False, True], ids=["every-campus", "one-campus"])
def test_the_manager_workbook_says_what_their_reports_page_says(api, manager, worked, nrb, narrowed):
    query = {"days": 30, **({"campus": nrb.pk} if narrowed else {})}
    api.force_authenticate(manager)
    page = api.get(reverse("analytics:overview"), query).json()

    _assert_workbook_matches(_workbook(api, manager, **query), page)


def _busy_campus(requester, nrb, electrical, nrb_electrician):
    """Live work for the Nairobi electrician, one job of it overdue, and an idle colleague."""
    factories.make_ticket(requester, nrb, electrical, status="in_progress", assigned_to=nrb_electrician)
    late = factories.make_ticket(requester, nrb, electrical, status="in_progress", assigned_to=nrb_electrician)
    type(late).objects.filter(pk=late.pk).update(resolution_due_at=timezone.now() - timedelta(hours=1))
    return factories.make_technician("idle_one", nrb, [electrical])


@pytest.mark.parametrize("role", ["nrb_hos", "nrb_hod"])
def test_a_supervisors_workbook_says_what_their_reports_page_says(
    api, request, worked, requester, nrb, electrical, nrb_electrician, role
):
    idle = _busy_campus(requester, nrb, electrical, nrb_electrician)
    user = request.getfixturevalue(role)
    api.force_authenticate(user)
    page = api.get(reverse("analytics:overview"), {"days": 30}).json()
    wb = _workbook(api, user, days=30)

    _assert_workbook_matches(wb, page)
    people = {r["Technician"]: r for r in _sheet(wb, "Technicians")}
    mine = next(t for t in page["technicians"] if t["key"] == nrb_electrician.pk)
    assert people[mine["label"]]["Open now"] == 2  # two live, one of them overdue
    assert idle.pk in {t["key"] for t in page["technicians"]}
    assert len(people) == len(page["technicians"]) == 2


@pytest.mark.parametrize("role", ["nrb_hos", "nrb_hod"])
def test_a_supervisors_pdf_prints_the_pages_technician_figures(
    api, monkeypatch, request, worked, requester, nrb, electrical, nrb_electrician, role
):
    _busy_campus(requester, nrb, electrical, nrb_electrician)
    user = request.getfixturevalue(role)
    api.force_authenticate(user)
    page = api.get(reverse("analytics:overview"), {"days": 30}).json()
    words = _pdf_words(api, user, monkeypatch, days=30)

    # The name column is a paragraph; the table's own strings follow its last heading, row by row.
    start = words.index("Time to resolve", words.index("Technicians")) + 2
    assert len(page["technicians"]) == 2
    for n, t in enumerate(page["technicians"]):
        row = words[start + 5 * n:start + 5 * (n + 1)]
        assert row == [
            str(t["total"]), str(t["open_now"]), _cell_text(t["resolution_sla_pct"]),
            _duration_text(t["resolution_seconds_avg"]), _rating_text(t["rating_avg"]),
        ]


def _duration_text(seconds):
    """formatSeconds in src/utils/date.ts."""
    if seconds is None:
        return "—"
    hours, rest = divmod(int(seconds), 3600)
    return f"{hours}h {rest // 60}m" if hours else f"{rest // 60}m"


def _rating_text(value):
    return "—" if value is None else f"{value:.2f}".rstrip("0").rstrip(".")


def _seconds(value):
    return None if value is None else timedelta(seconds=value)


def _assert_workbook_matches(wb, page):
    """Every figure on every sheet, against the page's own envelope."""
    current, before = page["current"], page["previous"]
    headline = {r["Figure"]: r for r in _sheet_from(wb, "Summary", "Figure") if r["Figure"]}
    assert headline["Raised"]["Value"] == current["total"]
    assert headline["Raised"]["Period before"] == before["total"]
    assert headline["Resolved"]["Value"] == current["settled"]
    assert headline["Resolved"]["Period before"] == before["settled"]
    assert headline["On hold now"]["Value"] == sum(b["total"] for b in page["bottlenecks"])
    assert headline["Resolved on time"]["Value"] == _pct(current["resolution_sla_pct"])
    assert headline["Resolved on time"]["Period before"] == _pct(before["resolution_sla_pct"])
    assert headline["Time to resolve"]["Value"] == _seconds(current["resolution_seconds_avg"])
    assert headline["Time to resolve"]["Period before"] == _seconds(before["resolution_seconds_avg"])
    assert "Open right now" not in headline
    assert headline["Rating"]["Value"] == current["rating_avg"]
    assert [(r["Reason"], r["On hold now"], r["By trade"]) for r in _sheet(wb, "Holds")] == [
        (b["label"], b["total"], " · ".join(f"{t['label']} {t['total']}" for t in b["trades"]) or None)
        for b in page["bottlenecks"]
    ]

    if "technicians" in page:
        assert _sheet(wb, "Technicians") == [
            {"Technician": t["label"], "Jobs": t["total"], "Open now": t["open_now"],
             "Resolved on time": _pct(t["resolution_sla_pct"]), "Time to resolve": _seconds(t["resolution_seconds_avg"]),
             "Rating": t["rating_avg"]}
            for t in page["technicians"]
        ]
        assert "Campuses" not in wb.sheetnames and "Trades by campus" not in wb.sheetnames
        page = {**page, "campuses": [], "trade_by_campus": None}
    else:
        assert "Technicians" not in wb.sheetnames

    campuses = {r["Campus"]: r for r in _sheet(wb, "Campuses")} if page["campuses"] else {}
    assert set(campuses) == {c["label"] for c in page["campuses"]}
    for c in page["campuses"]:
        assert campuses[c["label"]] == {
            "Campus": c["label"], "Raised": c["total"], "Open now": c["open_now"], "Overdue": c["overdue"],
            "With HOD": c["at_hod"], "Resolved on time": _pct(c["resolution_sla_pct"]),
            "Time to resolve": _seconds(c["resolution_seconds_avg"]), "Rating": c["rating_avg"],
            "Technicians": c["technicians"], "Jobs each": c["jobs_each"],
        }

    assert _sheet(wb, "Trades") == [
        {"Trade": t["label"], "Raised": t["total"], "Resolved on time": _pct(t["resolution_sla_pct"]), "Rating": t["rating_avg"]}
        for t in page["trades"]
    ]

    grid = page["trade_by_campus"]
    if grid:
        rows = {r["Campus"]: r for r in _sheet(wb, "Trades by campus")}
        for row in grid["rows"]:
            assert rows[row["campus"]] == {
                "Campus": row["campus"], **{t: row["counts"].get(t, 0) for t in grid["trades"]}, "All": row["total"]}
        assert rows["All"] == {
            "Campus": "All", **{t: grid["totals"].get(t, 0) for t in grid["trades"]}, "All": grid["total"]}

    assert [(r["Raised"], r["Settled"]) for r in _sheet(wb, "Tickets raised")] == [
        (p["total"], p["settled"]) for p in page["trend"]]
    assert [(r["Building"], r["Faults"], r["Open now"]) for r in _sheet(wb, "Facility Health")] == [
        (f["label"], f["total"], f["open_now"]) for f in page["facilities"]]


def _sheet_from(wb, name, heading):
    """A sheet's rows under the header row that starts with `heading`."""
    rows = list(wb[name].iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r and r[0] == heading)
    out = []
    for r in rows[start + 1:]:
        if not r or r[0] is None:
            break
        out.append(dict(zip(rows[start], r)))
    return out


def test_the_managers_pdf_prints_the_pages_campus_figures(api, monkeypatch, manager, worked):
    api.force_authenticate(manager)
    page = api.get(reverse("analytics:overview"), {"days": 30}).json()
    words = _pdf_words(api, manager, monkeypatch, days=30)

    # The name column is a paragraph; the table's own strings follow its last heading, row by row.
    start = words.index("Jobs each") + 1
    for n, c in enumerate(page["campuses"]):
        row = words[start + 9 * n:start + 9 * (n + 1)]
        assert row[:4] == [str(c["total"]), str(c["open_now"]), str(c["overdue"]), str(c["at_hod"])]
        assert row[7] == str(c["technicians"]) and row[4] == _cell_text(c["resolution_sla_pct"])


def _cell_text(pct):
    return "—" if pct is None else f"{pct:g}%"


def test_the_pdf_writes_durations_and_percentages_as_the_page_does():
    from apps.analytics.export import _cell

    # formatSeconds and formatPct in src/utils.
    assert (_cell(8100, "duration"), _cell(101_760, "duration"), _cell(900, "duration")) == (
        "2h 15m", "28h 16m", "15m",
    )
    assert (_cell(100.0, "pct"), _cell(72.3, "pct")) == ("100%", "72.3%")


def test_the_manager_workbook_lists_facility_health_and_attention(
    api, manager, requester, nrb, electrical, priorities
):
    block = factories.make_facility(nrb, "Wamalwa")
    socket = factories.make_service_item(electrical)
    for _ in range(3):
        factories.make_ticket(requester, nrb, electrical, service_item=socket, facility=block)

    wb = _workbook(api, manager, days=30)

    (row,) = _sheet(wb, "Facility Health")
    assert (row["Building"], row["Campus"], row["Faults"], row["By trade"]) == ("Wamalwa", "Nairobi", 3, "Electrical 3")
    assert row["Keeps coming back"] == "Faulty socket 3×"
    assert row["Suggests"] == "Electrical survey"
    attention = [r for r in wb["Summary"].iter_rows(values_only=True) if r and r[0] == "Facility"]
    assert attention and "Wamalwa" in attention[0][1]


def test_a_campus_narrows_the_managers_downloads(api, manager, worked, nrb):
    wb = _workbook(api, manager, campus=nrb.pk)

    assert _summary(wb)["Covers"] == "Nairobi campus"
    assert {t["Campus"] for t in _sheet(wb, "Tickets")} == {"Nairobi"}
    assert [r["Campus"] for r in _sheet(wb, "Campuses")] == ["Nairobi"]
    api.force_authenticate(manager)
    response = api.get(reverse(PDF), {"campus": nrb.pk})
    assert response.status_code == 200
    assert "nairobi-campus" in response["Content-Disposition"]


@pytest.mark.parametrize("bad", ["999999", "abc", "-1", ""])
def test_a_campus_that_is_not_one_is_refused_for_the_downloads(api, manager, worked, bad):
    api.force_authenticate(manager)
    assert api.get(reverse(XLSX), {"campus": bad}).status_code == 400


@pytest.mark.parametrize("role", ["nrb_hos", "nrb_hod", "nrb_electrician"])
def test_only_an_every_campus_report_takes_a_campus(api, request, worked, msa, role):
    api.force_authenticate(request.getfixturevalue(role))
    assert api.get(reverse(XLSX), {"campus": msa.pk}).status_code == 400


EVERY_CAMPUS_PDF = ["Headline figures", "Needs your attention", "Tickets raised", "Facility Health", "Campuses",
                    "Trades", "Trades by campus", "On hold now, by reason"]
CAMPUS_TABLE_COLUMNS = {"Raised", "Resolved", "Resolved on time", "With HOD", "Jobs each"}
ONE_CAMPUS_PDF = ["Headline figures", "Needs your attention", "Tickets raised", "Facility Health", "Technicians",
                  "Trades", "On hold now, by reason"]


@pytest.mark.parametrize(
    "role, sections, words_present, absent",
    [
        ("manager", EVERY_CAMPUS_PDF, CAMPUS_TABLE_COLUMNS, ("Trade performance",)),
        ("admin_user", EVERY_CAMPUS_PDF, CAMPUS_TABLE_COLUMNS, ("Trade performance",)),
        ("nrb_hos", ONE_CAMPUS_PDF, set(), ("Campuses", "Trades by campus", "Trade performance", "Where my jobs stand")),
        ("nrb_hod", ONE_CAMPUS_PDF, set(), ("Campuses", "Trades by campus", "Trade performance", "Where my jobs stand")),
    ],
)
def test_a_pdf_follows_its_reports_page(api, monkeypatch, request, worked, role, sections, words_present, absent):
    words = _pdf_words(api, request.getfixturevalue(role), monkeypatch)

    positions = [words.index(title) for title in sections]
    assert positions == sorted(positions)
    assert words_present <= set(words)
    for title in absent:
        assert title not in words


def _pdf_words(api, user, monkeypatch, **params):
    """Every heading, cell and legend label the PDF is built from."""
    from reportlab.graphics.charts import legends
    from reportlab import platypus

    words = []

    class Paragraph(platypus.Paragraph):
        def __init__(self, text, *args, **kwargs):
            words.append(text)
            super().__init__(text, *args, **kwargs)

    class Table(platypus.Table):
        def __init__(self, rows, *args, **kwargs):
            words.extend(cell for row in rows for cell in row if isinstance(cell, str))
            super().__init__(rows, *args, **kwargs)

    class Legend(legends.Legend):
        def __setattr__(self, name, value):
            if name == "colorNamePairs":
                words.extend(label for _, label in value)
            super().__setattr__(name, value)

    monkeypatch.setattr(platypus, "Paragraph", Paragraph)
    monkeypatch.setattr(platypus, "Table", Table)
    monkeypatch.setattr(legends, "Legend", Legend)
    api.force_authenticate(user)
    response = api.get(reverse(PDF), params)
    assert response.status_code == 200 and response.content.startswith(b"%PDF")
    return words


def test_a_one_trade_technicians_pdf_has_no_trade_sections(
    api, monkeypatch, nrb_electrician, nrb_hos, nrb_electrical_ticket
):
    assign_ticket(nrb_electrical_ticket, nrb_electrician, nrb_hos)
    words = _pdf_words(api, nrb_electrician, monkeypatch)

    assert "Where my jobs stand" in words and "Finished" in words
    for absent in ("Trade", "Trade performance", "Settled", "Avg resp."):
        assert absent not in words


def test_a_two_trade_technicians_pdf_keeps_its_trade_sections(
    api, monkeypatch, nrb, electrical, plumbing, nrb_hos, nrb_electrical_ticket, nrb_plumbing_ticket
):
    both = factories.make_technician("elec_and_plumb", nrb, [electrical, plumbing])
    assign_ticket(nrb_electrical_ticket, both, nrb_hos)
    assign_ticket(nrb_plumbing_ticket, both, nrb_hos)
    words = _pdf_words(api, both, monkeypatch)

    assert {"Where my jobs stand", "Trade", "Trade performance", "Finished"} <= set(words)
    # Response time is how fast work is assigned, which is the HOS's figure.
    assert "Avg resp." not in words
