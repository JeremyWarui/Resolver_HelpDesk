"""
The Excel export: one workbook, a Summary sheet and the tickets behind it.

The Summary is the same `metrics` calls the dashboard makes, over the same
scope and window as the Tickets sheet beside it. Nothing is recomputed here, so
the two can never disagree about a number.
"""

from datetime import timedelta
from io import BytesIO

import openpyxl
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.accounts.identity import display_name

from .metrics import bottlenecks, group, raised_in, resolve_window, totals

_HEADER_FILL = PatternFill("solid", fgColor="1E40AF")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LABEL_FILL = PatternFill("solid", fgColor="EFF6FF")
_SECTION_FONT = Font(bold=True)

_TIMEFRAME_DAYS = {"day": 1, "week": 7, "month": 30, "quarter": 90, "year": 365}

_TICKET_COLUMNS = [
    "Ticket No",
    "Status",
    "Priority",
    "Level",
    "Raised By",
    "Campus",
    "Trade",
    "Service Item",
    "Facility",
    "Assigned To",
    "Description",
    "Resolution Due",
    "Resolved At",
    "Closed At",
    "Created At",
    "Total Paused (mins)",
]


def workbook_window(params, qs) -> dict:
    """The window both sheets cover, from the report's own date vocabulary.

    `timeframe=all` is the case that needs care: it stretches back to the
    earliest ticket in scope, where `resolve_window` given nothing would quietly
    produce a thirty-day window instead.
    """
    start, end = params.get("start_date"), params.get("end_date")
    if start and end:
        return resolve_window({"date_from": start, "date_to": end})

    days = _TIMEFRAME_DAYS.get(params.get("timeframe", "all"))
    if days:
        return resolve_window({"days": days})

    now = timezone.now()
    earliest = qs.order_by("created_at").values_list("created_at", flat=True).first()
    return {
        "from": earliest or (now - timedelta(days=30)),
        "to": now,
        "bucket": "month",
    }


def _style_headers(ws, columns: int) -> None:
    ws.row_dimensions[1].height = 20
    for column in range(1, columns + 1):
        cell = ws.cell(row=1, column=column)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN


def _auto_width(ws) -> None:
    for cells in ws.columns:
        longest = max((len(str(c.value or "")) for c in cells), default=10)
        # Capped, so one long description cannot own the sheet.
        width = min(longest + 2, 50)
        ws.column_dimensions[get_column_letter(cells[0].column)].width = width


def _moment(dt) -> str:
    return dt.strftime("%Y-%m-%d %H:%M") if dt else ""


def _hours(seconds):
    return round(seconds / 3600, 1) if seconds is not None else ""


def _section(ws, title, rows) -> None:
    """A titled block of rows, with a blank line above it."""
    ws.append([])
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = _SECTION_FONT
    for row in rows:
        ws.append(list(row))


def _summary_sheet(ws, qs, window) -> None:
    figures = totals(qs, window)

    ws.append(["Metric", "Value"])
    _style_headers(ws, 2)
    for label, value in (
        ("Period from", _moment(window["from"])),
        ("Period to", _moment(window["to"])),
        ("Tickets raised", figures["total"]),
        ("Tickets settled", figures["settled"]),
        ("Open backlog", figures["open_backlog"]),
        ("Overdue now", figures["overdue"]),
        ("Missed target", figures["missed"]),
        ("Escalated", figures["escalated"]),
        ("Response SLA met (%)", figures["response_sla_pct"] or ""),
        ("Resolution SLA met (%)", figures["resolution_sla_pct"] or ""),
        ("Average response (hours)", _hours(figures["response_seconds_avg"])),
        ("Average resolution (hours)", _hours(figures["resolution_seconds_avg"])),
        ("Average rating", figures["rating_avg"] or ""),
        ("Ratings received", figures["rating_count"]),
    ):
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).fill = _LABEL_FILL

    for title, dimensions in (
        ("Tickets by status", ("status",)),
        ("Tickets by trade", ("trade",)),
        ("Tickets by technician", ("technician",)),
    ):
        _section(
            ws,
            title,
            [(row["label"], row["total"]) for row in group(qs, window, dimensions)],
        )

    # Facility health is the two-dimension cut: which building keeps breaking,
    # and which craft it keeps breaking for.
    _section(
        ws,
        "Facility health (facility by trade)",
        [
            (row["label"], row["label_2"], row["total"])
            for row in group(qs, window, ("facility", "trade"))
        ],
    )
    # A snapshot, not a period figure: the ticket column behind it is cleared
    # when work resumes. This sheet is where that snapshot becomes a record.
    _section(
        ws,
        "On hold now",
        [
            (row["label"], row["total"], _moment(row["since"]))
            for row in bottlenecks(qs)
        ],
    )

    _auto_width(ws)


def _tickets_sheet(ws, qs, window) -> None:
    ws.append(_TICKET_COLUMNS)
    _style_headers(ws, len(_TICKET_COLUMNS))
    ws.freeze_panes = "A2"

    tickets = raised_in(qs, window).select_related(
        "raised_by",
        "requester_campus",
        "trade",
        "service_item",
        "priority",
        "assigned_to",
        "facility",
    )
    for ticket in tickets:
        ws.append(
            [
                ticket.ticket_no,
                ticket.get_status_display(),
                ticket.priority.name if ticket.priority_id else "",
                ticket.get_current_level_display(),
                display_name(ticket.raised_by) if ticket.raised_by_id else "",
                ticket.requester_campus.name if ticket.requester_campus_id else "",
                ticket.trade.name if ticket.trade_id else "",
                ticket.service_item.name if ticket.service_item_id else "",
                ticket.facility.name,
                # Never blank: in a spreadsheet a blank reads as missing data.
                display_name(ticket.assigned_to) if ticket.assigned_to_id else "Unassigned",
                (ticket.description or "")[:200],
                _moment(ticket.resolution_due_at),
                _moment(ticket.resolved_at),
                _moment(ticket.closed_at),
                _moment(ticket.created_at),
                int(ticket.accumulated_pause.total_seconds() / 60),
            ]
        )

    _auto_width(ws)


def build_workbook(qs, window):
    """The whole workbook for one scope and window."""
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    _summary_sheet(summary, qs, window)
    _tickets_sheet(wb.create_sheet("Tickets"), qs, window)
    return wb


def workbook_response(wb) -> HttpResponse:
    """The workbook as a download."""
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = 'attachment; filename="tickets_report.xlsx"'
    return response
