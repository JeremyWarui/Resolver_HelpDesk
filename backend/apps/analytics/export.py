"""
The downloadable reports: an Excel workbook and a PDF summary.

Both are drawn from `report_data`, which reads every figure from `metrics` over
`reporting_ticket_qs`, so a download says exactly what the reports page says
for the same role and period. The manager's, the HOD's and the HOS's are laid
out as their Reports page and read `overview.build`, as it does. Times are
shown in Kenyan time.
"""

import io
import re
from datetime import timedelta
from decimal import ROUND_HALF_UP, Decimal
from xml.sax.saxutils import escape

from django.utils import timezone

from apps.accounts.identity import display_name
from apps.accounts.models import RoleAssignment
from apps.facilities.models import Facility
from apps.tickets.constants import LEVELS, PENDING_REASON_LABELS, STATUSES
from apps.tickets.scope import reporting_ticket_qs
from apps.tickets.sla import sla_outcome

from . import overview
from .metrics import REPORT_TZ, bottlenecks, group, measured, raised_in, totals, trend

# KSG brand colours for the reports. White on KSG Brown is 5.70:1.
KSG_BROWN = "7F622C"
KSG_BROWN_DEEP = "3E2C12"
KSG_GREEN_DARK = "6F7A00"

# Who may download which format.
XLSX_ROLES = ("admin", "manager", "hod", "hos", "technician")
PDF_ROLES = ("manager", "hod", "hos", "technician")

# "On hold", as the app says it, rather than the model's "Pending".
STATUS_LABELS = {**dict(STATUSES), "pending": "On hold"}
LEVEL_LABELS = dict(LEVELS)

# The columns of every breakdown sheet and table, in order: (key, heading).
BREAKDOWN_COLUMNS = [
    ("total", "Raised"),
    ("settled", "Settled"),
    ("open", "Open"),
    ("completion_pct", "Completion %"),
    ("response_seconds_avg", "Avg response"),
    ("resolution_seconds_avg", "Avg resolution"),
    ("resolution_sla_pct", "Resolution SLA %"),
    ("response_sla_pct", "Response SLA %"),
    ("missed", "Missed"),
    ("escalated", "Escalated"),
    ("rating_avg", "Rating"),
    ("rating_count", "Ratings"),
]


def scope_label(user, role, campus=None):
    """Who and what a report covers, said in words."""
    if campus:
        return f"{campus.name} campus"
    if role in overview.EVERY_CAMPUS_ROLES:
        return "Every campus"
    if role == "technician":
        return f"Jobs assigned to {display_name(user)}"
    assignment = RoleAssignment.objects.filter(user=user).select_related("campus").first()
    if assignment and assignment.campus:
        return f"{assignment.campus.name} campus"
    return "Your campus"


# Characters Excel refuses in a cell, which would fail the whole download.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# A cell starting with one of these is read by Excel as a formula.
_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")


def cell_text(value):
    """Text that is safe in a spreadsheet cell, from anything a person typed.

    Strips the control characters Excel rejects, and quotes a leading formula
    character so `=HYPERLINK(...)` in a comment stays text.
    """
    if not isinstance(value, str):
        return value
    value = _ILLEGAL.sub("", value)
    return "'" + value if value.startswith(_FORMULA_LEADS) else value


def local(dt):
    """A naive Kenyan-time datetime for a cell, or None."""
    return timezone.localtime(dt, REPORT_TZ).replace(tzinfo=None) if dt else None


def duration_text(seconds):
    """3h 07m, or 2d 4h for anything a day or longer."""
    if seconds is None:
        return "—"
    minutes = int(seconds // 60)
    days, rem = divmod(minutes, 24 * 60)
    hours, mins = divmod(rem, 60)
    if days:
        return f"{days}d {hours}h"
    return f"{hours}h {mins:02d}m"


def _facility_rows(qs, window):
    """Each facility with tickets in the window, and the trade it breaks for most."""
    by_pair = group(qs, window, ["facility", "trade"])
    top_trade = {}
    # Ties go to the trade first in the alphabet, so the same data always
    # names the same trade.
    for row in sorted(by_pair, key=lambda r: (-r["total"], r["label_2"] or "")):
        top_trade.setdefault(row["key"], (row["label_2"], row["total"]))
    rows = group(qs, window, ["facility"])
    # Every campus has an "Administration Block", so a row names its campus.
    campus_of = dict(
        Facility.objects.filter(pk__in=[r["key"] for r in rows]).values_list(
            "pk", "campus__name"
        )
    )
    for row in rows:
        row["top_trade"] = top_trade.get(row["key"], ("—", 0))[0]
        row["campus"] = campus_of.get(row["key"], "")
    return rows


def _ticket_rows(qs, window):
    """One row per ticket raised in the window, oldest first."""
    tickets = (
        measured(raised_in(qs, window))
        .select_related(
            "requester_campus", "trade", "service_item", "facility",
            "raised_by", "assigned_to", "priority",
        )
        .order_by("created_at")
    )
    rows = []
    for t in tickets:
        rows.append(
            {
                "ticket_no": t.ticket_no,
                "campus": t.requester_campus.name,
                "trade": t.trade.name,
                "service": t.service_item.name,
                "facility": t.facility.name,
                "where": t.location_detail,
                "raised_by": display_name(t.raised_by),
                "assigned_to": display_name(t.assigned_to) if t.assigned_to else "",
                "priority": t.priority.name,
                "status": STATUS_LABELS.get(t.status, t.status),
                "level": LEVEL_LABELS.get(t.current_level, t.current_level),
                "hold_reason": PENDING_REASON_LABELS.get(t.pending_reason, ""),
                "raised": local(t.created_at),
                "assigned": local(t.first_assigned_at),
                "resolved": local(t.resolved_at),
                "due": local(t.resolution_due_at),
                "sla": (sla_outcome(t) or "").capitalize(),
                "rating": t.rating,
                "rating_comment": t.rating_comment,
            }
        )
    return rows


def report_data(user, role, window, campus=None, campus_ids=None):
    """Everything both formats print, for this caller, period and campus.

    `campus_ids` are the campuses an HOS or HOD runs, for their technician rows.
    """
    qs = reporting_ticket_qs(user, role)
    if campus:
        qs = qs.filter(requester_campus=campus)
    heading = {
        "title": "Maintenance Helpdesk Report",
        "scope": scope_label(user, role, campus),
        "role": role,
        "from": local(window["from"]),
        "to": local(window["to"]),
        "generated": local(timezone.now()),
    }
    if role in overview.OVERVIEW_ROLES:
        return {
            **heading,
            "overview": overview.build(qs, window, role, campus_ids),
            "tickets": _ticket_rows(qs, window),
        }
    by_status = {row["key"]: row["total"] for row in group(qs, window, ["status"])}
    return {
        **heading,
        "totals": totals(qs, window),
        "by_status": {code: by_status.get(code, 0) for code, _ in STATUSES},
        "by_campus": group(qs, window, ["campus"]) if role in overview.EVERY_CAMPUS_ROLES else None,
        "by_trade": group(qs, window, ["trade"]),
        "by_technician": None if role == "technician" else group(qs, window, ["technician"]),
        "by_facility": _facility_rows(qs, window),
        "holds": bottlenecks(qs),
        "trend": trend(qs, window, window["bucket"]),
        "tickets": _ticket_rows(qs, window),
    }


def summary_rows(data):
    """The headline figures as (label, value, kind), kind naming the format.

    A technician's are about their own work: how many jobs, how fast, how many
    on time, and how they were rated. The rest are for whoever runs a section.
    """
    t = data["totals"]
    s = data["by_status"]
    if data["role"] == "technician":
        return [
            ("Jobs assigned to me", t["total"], "count"),
            ("Active now", t["open_backlog"], "count"),
            ("Overdue now", t["overdue"], "count"),
            ("Jobs finished", t["settled"], "count"),
            ("On time", t["resolution_sla_pct"], "pct"),
            ("Missed deadlines", t["missed"], "count"),
            ("Average time to fix", t["resolution_seconds_avg"], "duration"),
            ("Assigned", s["assigned"], "count"),
            ("In progress", s["in_progress"], "count"),
            ("On hold", s["pending"], "count"),
            ("Resolved", s["resolved"], "count"),
            ("Closed", s["closed"], "count"),
            ("Escalated, now with the HOS", t["escalated_hos"], "count"),
            ("Escalated, now with the HOD", t["escalated_hod"], "count"),
            ("Average rating (of 5)", t["rating_avg"], "number"),
            ("Ratings given", t["rating_count"], "count"),
        ]
    return [
        ("Tickets raised", t["total"], "count"),
        ("Open (not yet assigned)", s["open"], "count"),
        ("Assigned", s["assigned"], "count"),
        ("In progress", s["in_progress"], "count"),
        ("On hold", s["pending"], "count"),
        ("Resolved", s["resolved"], "count"),
        ("Closed", s["closed"], "count"),
        ("Completion rate", t["completion_pct"], "pct"),
        ("Open backlog (now)", t["open_backlog"], "count"),
        ("Overdue (now)", t["overdue"], "count"),
        ("Average response time", t["response_seconds_avg"], "duration"),
        ("Average resolution time", t["resolution_seconds_avg"], "duration"),
        ("Resolution SLA", t["resolution_sla_pct"], "pct"),
        ("Response SLA", t["response_sla_pct"], "pct"),
        ("Missed deadlines", t["missed"], "count"),
        ("Escalated, now with the HOS", t["escalated_hos"], "count"),
        ("Escalated, now with the HOD", t["escalated_hod"], "count"),
        ("Average rating (of 5)", t["rating_avg"], "number"),
        ("Ratings given", t["rating_count"], "count"),
        ("Closed tickets rated", t["rated_pct"], "pct"),
    ]


def report_headline(ov):
    """The Reports page's stat cards as (label, value, before, kind, note).

    `before` is the period before where the card shows it, else None.
    """
    now, before = ov["current"], ov["previous"]
    return [
        ("Resolved on time", now["resolution_sla_pct"], before["resolution_sla_pct"], "pct",
         "Settled work, by its deadline"),
        ("Average time to fix", now["resolution_seconds_avg"], before["resolution_seconds_avg"], "duration",
         "Holds left out"),
        ("Overdue now", now["overdue"], None, "count", f"{ov['at_hod_now']} live jobs sit with the HOD"),
        ("Open right now", now["open_backlog"], None, "count", f"{ov['unassigned_now']} waiting for a technician"),
        ("Rating", now["rating_avg"], None, "number", f"{now['rating_count']} ratings"),
    ]


INSIGHT_KINDS = {
    "campus": "Campus", "trade": "Trade", "technician": "Technician", "staffing": "Staffing",
    "facility": "Facility", "escalated": "Escalated",
}
# Each kind's (text, cell) colours, as the Reports page tints its first column.
INSIGHT_TONES = {
    "campus": ("#991B1B", "#FEF2F2"), "escalated": ("#991B1B", "#FEF2F2"),
    "trade": ("#92400E", "#FFFBEB"), "staffing": ("#92400E", "#FFFBEB"),
    "facility": ("#1E40AF", "#EFF6FF"), "technician": ("#5B21B6", "#F5F3FF"),
}
PLAIN_TONE = ("#374151", "#F9FAFB")


def facility_trades(row):
    """A Facility Health row's trades in words: Plumbing 5 · Carpentry 1."""
    return " · ".join(f"{t['label']} {t['total']}" for t in row["trades"])


def hours_minutes(seconds):
    """28h 16m, as the Reports page writes a time to fix."""
    if seconds is None:
        return "—"
    hours, rest = divmod(int(seconds), 3600)
    return f"{hours}h {rest // 60}m" if hours else f"{rest // 60}m"


def facility_pattern(row):
    """What keeps coming back at a building, or its average fix when nothing repeats."""
    if row["repeats"]:
        return ", ".join(f"{r['label']} {r['total']}×" for r in row["repeats"])
    return f"Avg fix {hours_minutes(row['resolution_seconds_avg'])}" if row["resolution_seconds_avg"] else "—"


def filename(data, extension):
    """e.g. helpdesk-report-nairobi-campus-2026-09-15-to-2026-10-15.xlsx

    ASCII letters and digits only, so the name sits in a plain header.
    """
    slug = "".join(
        c if c.isascii() and c.isalnum() else "-" for c in data["scope"].lower()
    ).strip("-") or "report"
    while "--" in slug:
        slug = slug.replace("--", "-")
    return (
        f"helpdesk-report-{slug}-{data['from']:%Y-%m-%d}-to-{data['to']:%Y-%m-%d}.{extension}"
    )


# ─── Excel ──────────────────────────────────────────────────────────────────

PCT_FORMAT = "0.0%"
DURATION_FORMAT = "[h]:mm"
DATE_FORMAT = "yyyy-mm-dd hh:mm"
KINDS = {
    "completion_pct": "pct", "resolution_sla_pct": "pct", "response_sla_pct": "pct",
    "response_seconds_avg": "duration", "resolution_seconds_avg": "duration",
}
TICKET_COLUMNS = [
    ("ticket_no", "Ticket"), ("campus", "Campus"), ("trade", "Trade"),
    ("service", "Service"), ("facility", "Facility"), ("where", "Where"),
    ("raised_by", "Raised by"), ("assigned_to", "Assigned to"),
    ("priority", "Priority"), ("status", "Status"), ("level", "Escalation level"),
    ("hold_reason", "Hold reason"), ("raised", "Raised"), ("assigned", "First assigned"),
    ("resolved", "Resolved"), ("due", "Resolution due"), ("sla", "SLA"),
    ("rating", "Rating"), ("rating_comment", "Rating comment"),
]


def _append(ws, values):
    """Add a row, every text value made safe for a cell."""
    ws.append([cell_text(v) for v in values])


def _header(ws, headings):
    from openpyxl.styles import Alignment, Font, PatternFill

    _append(ws, headings)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=KSG_BROWN)
        cell.alignment = Alignment(vertical="center")
    # An address, not `ws.cell(...)`: reading a cell creates the row.
    ws.freeze_panes = f"A{ws.max_row + 1}"


def _widths(ws, minimum=10, maximum=48):
    from openpyxl.utils import get_column_letter

    for column in ws.columns:
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(column[0].column)].width = max(
            minimum, min(maximum, longest + 2)
        )


def _put(cell, value, kind):
    if kind == "duration":
        cell.value = None if value is None else timedelta(seconds=value)
        cell.number_format = DURATION_FORMAT
    elif kind == "pct":
        # A fraction, so Excel's own percent arithmetic and formatting work.
        cell.value = None if value is None else value / 100
        cell.number_format = PCT_FORMAT
    else:
        cell.value = value


def _row(ws, values):
    """Add a row of (value, kind) pairs, each cell formatted for its kind."""
    r = ws.max_row + 1
    for i, (value, kind) in enumerate(values, start=1):
        _put(ws.cell(row=r, column=i), cell_text(value), kind)


def _cover(ws, data):
    from openpyxl.styles import Font

    _append(ws, [data["title"]])
    ws["A1"].font = Font(bold=True, size=14, color=KSG_BROWN_DEEP)
    _append(ws, ["Covers", data["scope"]])
    _append(ws, ["Period", f"{data['from']:%d %b %Y} to {data['to']:%d %b %Y}"])
    _append(ws, ["Generated", data["generated"]])
    ws["B4"].number_format = DATE_FORMAT
    _append(ws, ["Times are in Kenyan time (EAT)."])
    _append(ws, [])


def _holds_sheet(wb, holds):
    ws = wb.create_sheet("Holds")
    _header(ws, ["Reason", "On hold now", "Oldest hold since"])
    for row in holds:
        _append(ws, [row["label"], row["total"], local(row["since"])])
        ws.cell(row=ws.max_row, column=3).number_format = DATE_FORMAT
    _widths(ws)


def _tickets_sheet(wb, tickets):
    ws = wb.create_sheet("Tickets")
    _header(ws, [h for _, h in TICKET_COLUMNS])
    dated = {"raised", "assigned", "resolved", "due"}
    for row in tickets:
        _append(ws, [row[key] for key, _ in TICKET_COLUMNS])
        for i, (key, _) in enumerate(TICKET_COLUMNS, start=1):
            if key in dated:
                ws.cell(row=ws.max_row, column=i).number_format = DATE_FORMAT
    ws.auto_filter.ref = ws.dimensions
    _widths(ws, minimum=12)


def _saved(wb) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_workbook(data) -> bytes:
    """The workbook: a summary sheet, one sheet per breakdown, and every ticket."""
    if "overview" in data:
        return _overview_workbook(data)
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _cover(ws, data)
    _header(ws, ["Figure", "Value"])
    for label, value, kind in summary_rows(data):
        _append(ws, [label])
        _put(ws.cell(row=ws.max_row, column=2), value, kind)
    _append(ws, [])
    _append(ws, ["Durations read hours:minutes. Response is raised to first assigned; "
                 "resolution is raised to resolved, with time on hold left out."])
    _widths(ws)

    def breakdown_sheet(title, rows, first_heading):
        sheet = wb.create_sheet(title)
        _header(sheet, [first_heading] + [h for _, h in BREAKDOWN_COLUMNS])
        for row in rows:
            _append(sheet, [row["label"]])
            r = sheet.max_row
            for i, (key, _) in enumerate(BREAKDOWN_COLUMNS, start=2):
                _put(sheet.cell(row=r, column=i), row.get(key), KINDS.get(key, "number"))
        _widths(sheet)

    if data["by_campus"] is not None:
        breakdown_sheet("By campus", data["by_campus"], "Campus")
    breakdown_sheet("By trade", data["by_trade"], "Trade")
    if data["by_technician"] is not None:
        breakdown_sheet("By technician", data["by_technician"], "Technician")

    ws = wb.create_sheet("By facility")
    _header(ws, ["Facility", "Campus", "Raised", "Open", "Escalated", "Most frequent trade"])
    for row in data["by_facility"]:
        _append(ws, [
            row["label"], row["campus"], row["total"], row["open"], row["escalated"],
            row["top_trade"],
        ])
    _widths(ws)

    _holds_sheet(wb, data["holds"])
    _tickets_sheet(wb, data["tickets"])
    return _saved(wb)


def _overview_workbook(data) -> bytes:
    """The workbook of a Reports page built on `overview`: a sheet per section, then every ticket."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ov = data["overview"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    _cover(ws, data)
    _header(ws, ["Figure", "Value", "Period before", "Note"])
    for label, value, before, kind, note in report_headline(ov):
        _row(ws, [(label, "text"), (value, kind), (before, kind), (note, "text")])
    _append(ws, [])
    _append(ws, ["Needs your attention"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color=KSG_BROWN_DEEP)
    for item in ov["insights"]:
        _append(ws, [INSIGHT_KINDS.get(item["kind"], item["kind"]), item["text"]])
    if not ov["insights"]:
        _append(ws, ["Nothing is below target this period."])
    _widths(ws, maximum=90)

    ws = wb.create_sheet("Facility Health")
    _header(ws, ["Building", "Campus", "Type", "Faults", "By trade", "Keeps coming back", "Suggests", "Open now"])
    for row in ov["facilities"]:
        _append(ws, [row["label"], row["campus"], row["type"], row["total"], facility_trades(row),
                     facility_pattern(row), row["suggestion"], row["open_now"]])
    _widths(ws)

    if "campuses" in ov:
        ws = wb.create_sheet("Campuses")
        _header(ws, ["Campus", "Raised", "Open now", "Overdue", "With HOD", "On time", "Avg fix", "Rating",
                     "Technicians", "Jobs each"])
        for row in ov["campuses"]:
            _row(ws, [(row["label"], "text"), (row["total"], "count"), (row["open_now"], "count"),
                      (row["overdue"], "count"), (row["at_hod"], "count"), (row["resolution_sla_pct"], "pct"),
                      (row["resolution_seconds_avg"], "duration"), (row["rating_avg"], "number"),
                      (row["technicians"], "count"), (row["jobs_each"], "number")])
        _widths(ws)

    if "technicians" in ov:
        ws = wb.create_sheet("Technicians")
        _header(ws, ["Technician", "Jobs", "Open now", "On time", "Avg fix", "Rating"])
        for row in ov["technicians"]:
            _row(ws, [(row["label"], "text"), (row["total"], "count"), (row["open_now"], "count"),
                      (row["resolution_sla_pct"], "pct"), (row["resolution_seconds_avg"], "duration"),
                      (row["rating_avg"], "number")])
        _widths(ws)

    ws = wb.create_sheet("Trades")
    _header(ws, ["Trade", "Raised", "On time", "Rating"])
    for row in ov["trades"]:
        _row(ws, [(row["label"], "text"), (row["total"], "count"), (row["resolution_sla_pct"], "pct"),
                  (row["rating_avg"], "number")])
    _widths(ws)

    if "trade_by_campus" in ov:
        grid = ov["trade_by_campus"]
        ws = wb.create_sheet("Trades by campus")
        _header(ws, ["Campus", *grid["trades"], "All"])
        for row in grid["rows"]:
            _append(ws, [row["campus"], *(row["counts"].get(t, 0) for t in grid["trades"]), row["total"]])
        _append(ws, ["All", *(grid["totals"].get(t, 0) for t in grid["trades"]), grid["total"]])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        _widths(ws)

    ws = wb.create_sheet("Tickets raised")
    _header(ws, ["From", "Raised", "Settled"])
    for point in ov["trend"]:
        _append(ws, [local(point["bucket"]), point["total"], point["settled"]])
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd"
    _widths(ws, minimum=12)

    _holds_sheet(wb, ov["bottlenecks"])
    _tickets_sheet(wb, data["tickets"])
    return _saved(wb)


# ─── PDF ────────────────────────────────────────────────────────────────────


def _cell(value, kind):
    if value is None:
        return "—"
    if kind == "pct":
        return f"{value:.1f}%"
    if kind == "duration":
        return duration_text(value)
    if kind == "hm":
        return hours_minutes(value)
    if kind == "hours":
        # Rounded as the page's `toFixed(1)` rounds: half away from zero, on the exact value.
        return f"{Decimal(value / 3600).quantize(Decimal('0.1'), ROUND_HALF_UP)}h"
    if kind == "page_pct":
        return f"{value:g}%"
    if kind == "number":
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def build_pdf(data) -> bytes:
    """A few pages to read or share: headline figures, three charts, the tables."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    h1 = styles["Title"].clone("ksg-title", textColor=colors.HexColor(f"#{KSG_BROWN_DEEP}"))
    h2 = styles["Heading2"].clone("ksg-heading", textColor=colors.HexColor(f"#{KSG_BROWN_DEEP}"))
    small = styles["BodyText"].clone("small", fontSize=8, leading=10, textColor=colors.grey)
    cell_style = styles["BodyText"].clone("cell", fontSize=8, leading=9.5)
    brown = colors.HexColor(f"#{KSG_BROWN}")
    green_dark = colors.HexColor(f"#{KSG_GREEN_DARK}")
    palette = [colors.HexColor(c) for c in (
        "#2563EB", "#16A34A", "#F59E0B", "#DC2626", "#7C3AED", "#0891B2", "#DB2777", "#65A30D",
    )]
    page_width = A4[0] - 30 * mm

    def table(headings, rows, col_widths=None, text_columns=(0,)):
        """Numbers right-aligned; the columns named in `text_columns` left."""
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), brown),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FBF8EF")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2D8BF")),
            ("ALIGN", (0, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        style += [("ALIGN", (c, 1), (c, -1), "LEFT") for c in text_columns]
        # Names go in as paragraphs, so a long one wraps inside its column and
        # a "<" or "&" in it is text, not markup.
        rows = [
            [Paragraph(escape(str(v)), cell_style) if i in text_columns else v
             for i, v in enumerate(row)]
            for row in rows
        ]
        t = Table([headings] + rows, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle(style))
        return t

    def pie(items):
        items = [(label, value) for label, value in items if value]
        d = Drawing(page_width / 2, 62 * mm)
        if not items:
            return d
        p = Pie()
        p.x, p.y, p.width, p.height = 10, 20, 40 * mm, 40 * mm
        p.data = [v for _, v in items]
        p.labels = None
        for i in range(len(items)):
            p.slices[i].fillColor = palette[i % len(palette)]
            p.slices[i].strokeColor = colors.white
        d.add(p)
        legend = Legend()
        legend.x, legend.y = 50 * mm, 52 * mm
        legend.fontSize = 7
        legend.alignment = "right"
        # One column, so the legend stays inside its half of the page.
        legend.columnMaximum = len(items)
        legend.colorNamePairs = [
            (palette[i % len(palette)], f"{label} ({value})") for i, (label, value) in enumerate(items)
        ]
        d.add(legend)
        return d

    # A technician's report is about their own jobs, and says so in its own words.
    mine = data["role"] == "technician"
    finished = "Finished" if mine else "Settled"

    def bars(points):
        # Room under the axis for the slanted date labels, inside the drawing.
        d = Drawing(page_width, 62 * mm)
        if not points:
            return d
        chart = VerticalBarChart()
        chart.x, chart.y = 25, 12 * mm
        chart.width, chart.height = page_width - 40, 40 * mm
        chart.data = [[p["total"] for p in points], [p["settled"] for p in points]]
        chart.bars[0].fillColor = brown
        chart.bars[1].fillColor = green_dark
        chart.categoryAxis.categoryNames = [
            f"{timezone.localtime(p['bucket'], REPORT_TZ):%d %b}" for p in points
        ]
        chart.categoryAxis.labels.fontSize = 6
        chart.categoryAxis.labels.angle = 45 if len(points) > 12 else 0
        chart.categoryAxis.labels.boxAnchor = "ne" if len(points) > 12 else "n"
        chart.valueAxis.valueMin = 0
        chart.valueAxis.labels.fontSize = 7
        chart.valueAxis.valueStep = max(1, (max(chart.data[0]) // 5) or 1)
        d.add(chart)
        legend = Legend()
        legend.x, legend.y = page_width - 70, 62 * mm
        legend.fontSize = 7
        legend.colorNamePairs = [(brown, "Raised"), (green_dark, finished)]
        d.add(legend)
        return d

    def breakdown(title, rows, first_heading, limit=None):
        columns = [
            ("total", "Raised", "count"), ("open", "Open", "count"),
            ("completion_pct", "Done", "pct"),
        ] + ([] if mine else [("response_seconds_avg", "Avg resp.", "duration")]) + [
            ("resolution_seconds_avg", "Avg resol.", "duration"),
            ("resolution_sla_pct", "SLA", "pct"), ("escalated", "Escal.", "count"),
            ("rating_avg", "Rating", "number"),
        ]
        rows = rows[:limit] if limit else rows
        body_rows = [
            [row["label"]] + [_cell(row.get(key), kind) for key, _, kind in columns]
            for row in rows
        ]
        widths_ = [page_width * 0.28] + [page_width * 0.09] * len(columns)
        return KeepTogether([
            Paragraph(title, h2),
            table([first_heading] + [h for _, h, _ in columns], body_rows or [["No tickets"] + [""] * len(columns)], widths_),
            Spacer(1, 4 * mm),
        ])

    def overview_story():
        """The Reports page built on `overview`, section by section."""
        ov = data["overview"]
        story = cover() + [Paragraph("Headline figures", h2)]
        cards = []
        for label, value, before, kind, note in report_headline(ov):
            kind = "hm" if kind == "duration" else kind
            kind = "page_pct" if kind == "pct" else kind
            shown = f"{_cell(value, kind)} / 5" if label == "Rating" and value is not None else _cell(value, kind)
            beside = f"{_cell(before, kind)} the period before" if before is not None else note
            cards.append([label, shown, beside])
        story += [
            table(["Figure", "Value", "Note"], cards, [page_width * w for w in (0.3, 0.2, 0.5)], text_columns=(0, 2)),
            Spacer(1, 4 * mm),
            Paragraph("Needs your attention", h2),
        ]
        if ov["insights"]:
            # A bordered box ruled inside, with no header row, as the page sets it out.
            attention = Table(
                [[Paragraph(f"<font color='{INSIGHT_TONES.get(i['kind'], PLAIN_TONE)[0]}'><b>"
                            f"{INSIGHT_KINDS.get(i['kind'], i['kind'])}</b></font>", cell_style),
                  Paragraph(escape(i["text"]), cell_style)] for i in ov["insights"]],
                colWidths=[page_width * 0.16, page_width * 0.84],
            )
            attention.setStyle(TableStyle([
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D3C3")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E8E3D6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                *(("BACKGROUND", (0, n), (0, n), colors.HexColor(INSIGHT_TONES.get(i["kind"], PLAIN_TONE)[1]))
                  for n, i in enumerate(ov["insights"])),
            ]))
            story.append(attention)
        else:
            story.append(Paragraph("Nothing is below target this period.", body))

        health = [
            [f"{row['label']} ({row['campus']} · {row['type']})", str(row["total"]), facility_trades(row),
             facility_pattern(row), row["suggestion"], str(row["open_now"])]
            for row in ov["facilities"]
        ]
        story += [Spacer(1, 4 * mm), KeepTogether([
            Paragraph("Facility Health", h2),
            Paragraph("Buildings where faults keep coming back this period, and in which trade.", small),
            table(["Building", "Faults", "By trade", "Keeps coming back", "Suggests", "Open"],
                  health or [["No building shows a pattern this period.", "", "", "", "", ""]],
                  [page_width * w for w in (0.24, 0.08, 0.22, 0.18, 0.2, 0.08)], text_columns=(0, 2, 3, 4)),
        ])]

        target = ov["target_pct"]
        if "campuses" in ov:
            campuses = [
                [row["label"], str(row["total"]), str(row["open_now"]), str(row["overdue"]), str(row["at_hod"]),
                 _cell(row["resolution_sla_pct"], "page_pct"), _cell(row["resolution_seconds_avg"], "hours"),
                 _cell(row["rating_avg"], "number"), str(row["technicians"]), _cell(row["jobs_each"], "number")]
                for row in ov["campuses"]
            ]
            story += [Spacer(1, 4 * mm), KeepTogether([
                Paragraph("Campuses", h2),
                Paragraph(f"Target: {target}% on time.", small),
                table(["Campus", "Raised", "Open now", "Overdue", "With HOD", "On time", "Avg fix", "Rating",
                       "Technicians", "Jobs each"],
                      campuses or [["No tickets"] + [""] * 9],
                      [page_width * 0.13] + [page_width * 0.09] * 7 + [page_width * 0.12] * 2),
            ])]

        if "technicians" in ov:
            people = [
                [row["label"], str(row["total"]), str(row["open_now"]),
                 _cell(row["resolution_sla_pct"], "page_pct"), _cell(row["resolution_seconds_avg"], "hours"),
                 _cell(row["rating_avg"], "number")]
                for row in ov["technicians"]
            ]
            story += [Spacer(1, 4 * mm), KeepTogether([
                Paragraph("Technicians", h2),
                Paragraph(f"Jobs raised this period and assigned to each. Target: {target}% on time.", small),
                table(["Technician", "Jobs", "Open now", "On time", "Avg fix", "Rating"],
                      people or [["No technicians", "", "", "", "", ""]],
                      [page_width * 0.35] + [page_width * 0.13] * 5),
            ])]

        trades = [
            [row["label"], str(row["total"]), _cell(row["resolution_sla_pct"], "page_pct"), _cell(row["rating_avg"], "number")]
            for row in ov["trades"]
        ]
        story += [Spacer(1, 4 * mm), KeepTogether([
            Paragraph("Trades", h2),
            table(["Trade", "Raised", "On time", "Rating"], trades or [["No tickets", "", "", ""]],
                  [page_width * w for w in (0.4, 0.2, 0.2, 0.2)]),
        ])]

        if "trade_by_campus" in ov:
            grid = ov["trade_by_campus"]
            cells = [[row["campus"], *(str(row["counts"].get(t, 0)) for t in grid["trades"]), str(row["total"])]
                     for row in grid["rows"]]
            cells.append(["All", *(str(grid["totals"].get(t, 0)) for t in grid["trades"]), str(grid["total"])])
            share = 0.8 / (len(grid["trades"]) + 1)
            story += [Spacer(1, 4 * mm), KeepTogether([
                Paragraph("Trades by campus", h2),
                Paragraph("Tickets raised in each trade at each campus.", small),
                table(["Campus", *grid["trades"], "All"], cells,
                      [page_width * 0.2] + [page_width * share] * (len(grid["trades"]) + 1)),
            ])]

        story += [Spacer(1, 4 * mm), KeepTogether([Paragraph("Tickets raised", h2), bars(ov["trend"])])]

        holds = [
            [row["label"], str(row["total"]), f"{local(row['since']):%d %b %Y}" if row["since"] else "—"]
            for row in ov["bottlenecks"]
        ]
        story.append(KeepTogether([
            Paragraph("On hold now, by reason", h2),
            table(["Reason", "Tickets", "Oldest hold since"], holds or [["Nothing is on hold", "", ""]],
                  [page_width * 0.6, page_width * 0.15, page_width * 0.25]),
        ]))
        return story

    def cover():
        return [
            Paragraph(data["title"], h1),
            Paragraph(f"<b>Covers:</b> {escape(data['scope'])}", body),
            Paragraph(
                f"<b>Period:</b> {data['from']:%d %b %Y} to {data['to']:%d %b %Y}"
                f" &nbsp;&nbsp; <b>Generated:</b> {data['generated']:%d %b %Y %H:%M} EAT",
                body,
            ),
            Spacer(1, 5 * mm),
        ]

    def section_story():
        """Every other role's summary: headline, facilities, charts, then the tables."""
        story = cover() + [Paragraph("Headline figures", h2)]

        figures = summary_rows(data)
        half = (len(figures) + 1) // 2
        pairs = []
        for left, right in zip(figures[:half], figures[half:] + [None]):
            pairs.append(
                [left[0], _cell(left[1], left[2])]
                + ([right[0], _cell(right[1], right[2])] if right else ["", ""])
            )
        headline = table(["Figure", "Value", "Figure", "Value"], pairs,
                         [page_width * 0.32, page_width * 0.18] * 2, text_columns=(0, 2))
        story += [headline, Spacer(1, 2 * mm), Paragraph(
            (
                "Counts are of your jobs raised in the period, except active and overdue, which "
                "are as of now. Time to fix leaves out time on hold. On time is the share of "
                "finished jobs done by their deadline."
                if mine else
                "Counts are of tickets raised in the period, except backlog and overdue, which are "
                "as of now. Resolution time leaves out time on hold. SLA is the share of settled "
                "tickets finished by their deadline."
            ), small), Spacer(1, 4 * mm)]
        # A trade with no jobs in the period is no slice, so this counts trades worked, not held.
        several_trades = not mine or len(data["by_trade"]) > 1

        # Where the work is, straight after how much of it there is.
        hotspots = [
            [row["label"], row["campus"], str(row["total"]), str(row["open"]),
             str(row["escalated"]), row["top_trade"]]
            for row in data["by_facility"][:10]
        ]
        story.append(KeepTogether([
            Paragraph("Where your jobs were (top 10)" if mine else "Facility hotspots (top 10)", h2),
            table(["Facility", "Campus", "Raised", "Open", "Escalated", "Most frequent trade"],
                  hotspots or [["No tickets", "", "", "", "", ""]],
                  [page_width * w for w in (0.28, 0.16, 0.1, 0.1, 0.12, 0.24)],
                  text_columns=(0, 1, 5)),
        ]))

        story += [
            Spacer(1, 4 * mm),
            KeepTogether([
                Paragraph("Jobs raised and finished" if mine else "Tickets raised and settled", h2),
                bars(data["trend"]),
            ]),
            Spacer(1, 4 * mm),
        ]
        status_items = [(STATUS_LABELS[code], n) for code, n in data["by_status"].items()]
        trade_items = [(row["label"], row["total"]) for row in data["by_trade"]]
        if several_trades:
            pies = Table([[Paragraph("Where my jobs stand" if mine else "Status", body),
                           Paragraph("Trade", body)],
                          [pie(status_items), pie(trade_items)]],
                         colWidths=[page_width / 2] * 2)
        else:
            pies = Table([[Paragraph("Where my jobs stand", body)], [pie(status_items)]],
                         colWidths=[page_width])
        story += [KeepTogether([pies]), Spacer(1, 4 * mm)]

        if data["by_campus"] is not None:
            story.append(breakdown("Campus comparison", data["by_campus"], "Campus"))
        if several_trades:
            story.append(breakdown("Trade performance", data["by_trade"], "Trade"))
        if data["by_technician"] is not None:
            story.append(breakdown("Technician performance", data["by_technician"], "Technician"))

        holds = [
            [row["label"], str(row["total"]), f"{local(row['since']):%d %b %Y}" if row["since"] else "—"]
            for row in data["holds"]
        ]
        story.append(KeepTogether([
            Paragraph("On hold now, by reason", h2),
            table(["Reason", "Tickets", "Oldest hold since"], holds or [["Nothing is on hold", "", ""]],
                  [page_width * 0.6, page_width * 0.15, page_width * 0.25]),
        ]))
        return story

    story = overview_story() if "overview" in data else section_story()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(15 * mm, 10 * mm, f"{data['title']} · {data['scope']}")
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=18 * mm,
        title=data["title"], author="Resolver", subject=data["scope"],
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
